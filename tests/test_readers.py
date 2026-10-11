import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from staffbot.readers import (Dataset, InputError, ParseOptions, Record, generic_text,
                             load_file, split_whitespace, text_records)


class ReaderTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.addCleanup(self.temp.cleanup)

    def file(self, name, text):
        path = self.root / name
        path.write_text(text, encoding="utf-8")
        return path

    def test_quoted_names_apostrophes_paths_and_escaped_quotes(self):
        row = split_whitespace('001 "Mary Jane" O\'Neil C:\\staff "A ""quoted"" name" ""', "Line 1")
        self.assertEqual(row, ("001", "Mary Jane", "O'Neil", "C:\\staff", 'A "quoted" name', ""))

    def test_plain_and_multiple_spaces(self):
        self.assertEqual(split_whitespace("  a   b\tc  ", "Line 1"), ("a", "b", "c"))

    def test_unclosed_quote_is_not_silently_split(self):
        with self.assertRaisesRegex(InputError, "unclosed"):
            split_whitespace('1 "Mary Jane', "Line 2")

    def test_missing_space_after_quote_is_reported(self):
        with self.assertRaisesRegex(InputError, "space after"):
            split_whitespace('"Mary"Jane', "Line 2")

    def test_csv_quote_delimiter_and_blank_cells(self):
        records = text_records('001,"Jane, Mary",,HR\n', ParseOptions(separator="comma"))
        self.assertEqual(records[0].values, ("001", "Jane, Mary", "", "HR"))

    def test_pipe_missing_and_final_empty_fields_survive(self):
        records = text_records("001||HR|\n", ParseOptions(separator="pipe"))
        self.assertEqual(records[0].values, ("001", "", "HR", ""))

    def test_custom_separator(self):
        rows = text_records('001;"Mary Jane";HR', ParseOptions(separator="custom", custom_separator=";"))
        self.assertEqual(rows[0].values, ("001", "Mary Jane", "HR"))
        with self.assertRaises(InputError):
            text_records("a::b", ParseOptions(separator="custom", custom_separator="::"))

    def test_skip_header_and_reorder_columns(self):
        path = self.file("data.csv", "id,name,dept\n001,Mary Jane,HR\n")
        data = load_file(path, ParseOptions(separator="comma", skip_header=True, columns=(2, 1)))
        self.assertEqual(data.records[0].values, ("Mary Jane", "001"))
        data.validate()

    def test_column_out_of_range_is_reported(self):
        path = self.file("data.txt", "a b\n")
        with self.assertRaisesRegex(InputError, "outside"):
            load_file(path, ParseOptions(columns=(3,)))

    def test_utf16_notepad_file(self):
        path = self.root / "notepad.txt"
        path.write_bytes('001 "José García" HR'.encode("utf-16"))
        self.assertEqual(load_file(path).records[0].values[1], "José García")

    def test_encoding_is_not_guessed_after_utf8_fails(self):
        path = self.root / "old.txt"
        path.write_bytes("001 José HR".encode("cp1252"))
        with self.assertRaisesRegex(InputError, "Choose its encoding"):
            load_file(path)
        self.assertEqual(load_file(path, ParseOptions(encoding="cp1252")).records[0].values[1], "José")

    def test_inconsistent_rows_block_automation(self):
        data = Dataset([Record("Line 1", ("a", "b")), Record("Line 2", ("c",))])
        with self.assertRaisesRegex(InputError, "Line 2 has 1 fields"):
            data.validate()

    def test_embedded_control_characters_block_form_navigation(self):
        for value in ("a\nb", "a\tb", "a\x00b"):
            with self.assertRaises(InputError):
                Dataset([Record("Row 1", (value,))]).validate()

    def test_empty_file_and_doc_extension_reported(self):
        with self.assertRaises(InputError):
            load_file(self.file("empty.txt", " \n\n")).validate()
        with self.assertRaisesRegex(InputError, "Save it as .docx"):
            load_file(self.file("old.doc", ""))

    def test_generic_conversion_keeps_ids_and_known_zero_mask(self):
        self.assertEqual(generic_text("001"), "001")
        self.assertEqual(generic_text(1, "000"), "001")
        self.assertEqual(generic_text(1.0), "1")
        self.assertEqual(generic_text(date(2026, 10, 10)), "2026-10-10")
        self.assertEqual(generic_text(False), "FALSE")
        self.assertEqual(generic_text(None), "")
        with self.assertRaises(InputError):
            generic_text(float("nan"))

    def test_real_word_paragraphs_and_tables_keep_order(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('001 "Mary Jane" HR')
        table = doc.add_table(rows=1, cols=3)
        for cell, value in zip(table.rows[0].cells, ("002", "José García", "IT")):
            cell.text = value
        doc.add_paragraph('003 "John O\'Neil" Ops')
        path = self.root / "staff.docx"
        doc.save(path)
        data = load_file(path)
        data.validate()
        self.assertEqual([r.values[0] for r in data.records], ["001", "002", "003"])
        self.assertEqual(data.records[1].values[1], "José García")

    def test_merged_word_cells_are_rejected(self):
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "001"
        path = self.root / "merged.docx"
        doc.save(path)
        with self.assertRaisesRegex(InputError, "unmerge"):
            load_file(path)

    def test_real_excel_preserves_blank_fields_and_zero_ids(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = "Staff"
        sheet.append(["ID", "Name", "Dept", "Email"])
        sheet.append([1, "Mary Jane", None, "mary@example.com"])
        sheet["A2"].number_format = "000"
        sheet.append([None, "José García", "IT", None])
        path = self.root / "staff.xlsx"
        book.save(path)
        data = load_file(path, ParseOptions(skip_header=True, sheet="Staff"))
        data.validate()
        self.assertEqual(data.records[0].values, ("001", "Mary Jane", "", "mary@example.com"))
        self.assertEqual(data.records[1].values, ("", "José García", "IT", ""))
        self.assertEqual(data.records[1].source, "Staff row 3")

    def test_excel_uncalculated_formula_is_not_sent_as_blank(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(["001", "=1+2"])
        path = self.root / "formula.xlsx"
        book.save(path)
        with self.assertRaisesRegex(InputError, "formula has no saved result"):
            load_file(path)

    def test_excel_missing_sheet_is_reported(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.append([1, 2])
        path = self.root / "data.xlsx"
        book.save(path)
        with self.assertRaisesRegex(InputError, "Unknown sheet"):
            load_file(path, ParseOptions(sheet="Missing"))


if __name__ == "__main__":
    unittest.main()
