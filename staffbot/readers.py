"""Read unrelated source formats into the same representation: lists of strings."""
from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path


class InputError(ValueError):
    pass


@dataclass(frozen=True)
class ParseOptions:
    separator: str = "whitespace"
    custom_separator: str = "|"
    encoding: str = "auto"
    skip_header: bool = False
    sheet: str = ""
    # 1-based indexes; can select and reorder fields. Empty means every column.
    columns: tuple[int, ...] = ()


@dataclass(frozen=True)
class Record:
    source: str
    values: tuple[str, ...]


@dataclass
class Dataset:
    records: list[Record]
    warnings: list[str] = field(default_factory=list)

    @property
    def width(self) -> int:
        return max((len(r.values) for r in self.records), default=0)

    def validate(self) -> None:
        if not self.records:
            raise InputError("No data rows found. Check the file and header setting.")
        expected = len(self.records[0].values)
        if not expected:
            raise InputError("The first record contains no fields.")
        for record in self.records:
            if len(record.values) != expected:
                raise InputError(
                    f"{record.source} has {len(record.values)} fields; expected {expected}. "
                    'Quote values containing spaces, for example "Mary Jane", '
                    "or use a comma, tab, or pipe separator."
                )
            for value in record.values:
                if any(ch in value for ch in "\r\n\t\x00"):
                    raise InputError(
                        f"{record.source} contains a newline, tab, or NUL inside a field. "
                        "Replace it with a space before entering this record."
                    )


def generic_text(value, number_format: str = "General") -> str:
    """Keep source strings; stringify other Excel values without schema guessing."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise InputError("An Excel cell contains an infinite or NaN number.")
    if isinstance(value, (float, int)) and int(value) == value:
        text = str(int(value))
        # Common employee-ID format: 00000. Other display formats are not inferred.
        if re.fullmatch(r"0+", number_format or ""):
            return text.zfill(len(number_format))
        return text
    return str(value)


def read_text(path: Path, encoding: str) -> tuple[str, str]:
    raw = path.read_bytes()
    if encoding != "auto":
        try:
            return raw.decode(encoding), encoding
        except (UnicodeError, LookupError) as exc:
            raise InputError(f"Could not decode this file as {encoding}: {exc}") from exc
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), "utf-16"
    try:
        return raw.decode("utf-8-sig"), "utf-8"
    except UnicodeDecodeError as exc:
        raise InputError(
            "This file is not UTF-8 or BOM-marked UTF-16. "
            "Choose its encoding (for example cp1252), or save it as UTF-8 in Notepad."
        ) from exc


def split_whitespace(line: str, source: str) -> tuple[str, ...]:
    """Whitespace delimiters with quoted fields; ordinary apostrophes/backslashes survive."""
    fields: list[str] = []
    i, n = 0, len(line)
    while i < n:
        while i < n and line[i].isspace():
            i += 1
        if i == n:
            break
        if line[i] == '"':
            i += 1
            part: list[str] = []
            while i < n:
                if line[i] == '"':
                    if i + 1 < n and line[i + 1] == '"':
                        part.append('"')
                        i += 2
                        continue
                    i += 1
                    break
                part.append(line[i])
                i += 1
            else:
                raise InputError(f"{source}: unclosed double quote.")
            if i < n and not line[i].isspace():
                raise InputError(f"{source}: put a space after a closing double quote.")
            fields.append("".join(part))
        else:
            start = i
            while i < n and not line[i].isspace():
                i += 1
            fields.append(line[start:i])
    return tuple(fields)


def text_records(text: str, options: ParseOptions, prefix: str = "Line") -> list[Record]:
    if options.separator == "whitespace":
        return [Record(f"{prefix} {i}", split_whitespace(line, f"{prefix} {i}"))
                for i, line in enumerate(text.splitlines(), 1) if line.strip()]
    delimiters = {"tab": "\t", "comma": ",", "semicolon": ";", "pipe": "|"}
    delimiter = delimiters.get(options.separator, options.custom_separator)
    if len(delimiter) != 1 or delimiter in '\r\n\x00"':
        raise InputError("The custom separator must be one character, such as | or ;.")
    reader = csv.reader(io.StringIO(text), delimiter=delimiter, skipinitialspace=True, strict=True)
    result: list[Record] = []
    try:
        for values in reader:
            if values and any(v.strip() for v in values):
                result.append(Record(f"{prefix} {reader.line_num}", tuple(values)))
    except csv.Error as exc:
        raise InputError(f"Invalid delimited text near {prefix} {reader.line_num}: {exc}") from exc
    return result


def _dependency(module: str, package: str):
    import importlib
    try:
        return importlib.import_module(module)
    except ImportError as exc:
        raise InputError(f"Missing package {package}. Run the setup script first.") from exc


def _docx(path: Path, options: ParseOptions) -> Dataset:
    docx = _dependency("docx", "python-docx")
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = docx.Document(path)
    rows: list[Record] = []
    warnings = ["Word headers, footers, text boxes, pictures, and nested tables are not imported."]
    # iter_inner_content keeps paragraphs and tables in their original body order.
    for i, block in enumerate(document.iter_inner_content(), 1):
        if isinstance(block, Paragraph):
            rows.extend(text_records(block.text, options, f"Paragraph {i}, line"))
        elif isinstance(block, Table):
            for j, row in enumerate(block.rows, 1):
                if row.grid_cols_before or row.grid_cols_after:
                    raise InputError(f"Table {i}, row {j}: use a regular rectangular Word table.")
                cells = row.cells
                if len({id(cell._tc) for cell in cells}) != len(cells) or any(
                    cell._tc.vMerge is not None for cell in cells
                ):
                    raise InputError(f"Table {i}, row {j}: unmerge table cells before import.")
                values = tuple(cell.text for cell in cells)
                if any(v.strip() for v in values):
                    rows.append(Record(f"Table {i}, row {j}", values))
    return Dataset(rows, warnings)


def _trim_excel(rows: list[Record]) -> list[Record]:
    width = max((i + 1 for r in rows for i, v in enumerate(r.values) if v != ""), default=0)
    return [Record(r.source, r.values[:width]) for r in rows]


def _xlsx(path: Path, options: ParseOptions) -> Dataset:
    openpyxl = _dependency("openpyxl", "openpyxl")
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached = None
    warnings = ["Excel dates use ISO text. Currency, percentage, and locale display formats are not reproduced."]
    try:
        name = options.sheet or formulas.sheetnames[0]
        if name not in formulas.sheetnames:
            raise InputError(f"Unknown sheet {name!r}. Available: {', '.join(formulas.sheetnames)}")
        sheet = formulas[name]
        if sheet.max_row * sheet.max_column > 2_000_000:
            raise InputError("The worksheet range is too large. Copy just your data into a new sheet.")
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows: list[Record] = []
        for row_number, (row, cache_row) in enumerate(zip(sheet.iter_rows(), cached[name].iter_rows()), 1):
            values: list[str] = []
            for cell, cache_cell in zip(row, cache_row):
                value = cell.value
                if cell.data_type == "f":
                    value = cache_cell.value
                    if value is None:
                        raise InputError(
                            f"{name}!{cell.coordinate}: formula has no saved result. "
                            "Recalculate and save in Excel/LibreOffice, or paste values into a new file."
                        )
                if cell.data_type == "e" or cache_cell.data_type == "e":
                    raise InputError(f"{name}!{cell.coordinate}: fix the Excel error before importing.")
                values.append(generic_text(value, cell.number_format))
            if any(v != "" for v in values):
                rows.append(Record(f"{name} row {row_number}", tuple(values)))
        return Dataset(_trim_excel(rows), warnings)
    finally:
        formulas.close()
        if cached is not None:
            cached.close()


def _xls(path: Path, options: ParseOptions) -> Dataset:
    xlrd = _dependency("xlrd", "xlrd")
    book = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True)
    try:
        name = options.sheet or book.sheet_names()[0]
        if name not in book.sheet_names():
            raise InputError(f"Unknown sheet {name!r}. Available: {', '.join(book.sheet_names())}")
        sheet = book.sheet_by_name(name)
        rows: list[Record] = []
        for r in range(sheet.nrows):
            values: list[str] = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    value = ""
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    dt = xlrd.xldate_as_datetime(cell.value, book.datemode)
                    value = dt.date().isoformat() if dt.time() == time() else dt.isoformat(sep=" ")
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = generic_text(bool(cell.value))
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    raise InputError(f"{name} row {r + 1}, column {c + 1}: Excel error.")
                else:
                    xf = book.xf_list[sheet.cell_xf_index(r, c)]
                    fmt = book.format_map.get(xf.format_key)
                    value = generic_text(cell.value, fmt.format_str if fmt else "General")
                values.append(value)
            if any(v != "" for v in values):
                rows.append(Record(f"{name} row {r + 1}", tuple(values)))
        return Dataset(_trim_excel(rows), ["XLS uses saved cell values; formula freshness cannot be checked."])
    finally:
        book.release_resources()


def load_file(path: str | Path, options: ParseOptions = ParseOptions()) -> Dataset:
    path = Path(path).expanduser()
    if not path.is_file():
        raise InputError(f"File not found: {path}")
    extension = path.suffix.lower()
    if extension in (".txt", ".csv", ".tsv"):
        text, encoding = read_text(path, options.encoding)
        dataset = Dataset(text_records(text, options), [f"Text encoding: {encoding}."])
    elif extension == ".docx":
        dataset = _docx(path, options)
    elif extension in (".xlsx", ".xlsm"):
        dataset = _xlsx(path, options)
    elif extension == ".xls":
        dataset = _xls(path, options)
    elif extension == ".doc":
        raise InputError("Legacy .doc is not supported. Save it as .docx in Word or LibreOffice first.")
    else:
        raise InputError("Choose a TXT, CSV, TSV, DOCX, XLSX, XLSM, or XLS file.")
    if options.skip_header and dataset.records:
        dataset.records = dataset.records[1:]
    if options.columns:
        if any(c < 1 for c in options.columns):
            raise InputError("Column indexes start at 1.")
        selected: list[Record] = []
        for record in dataset.records:
            if max(options.columns) > len(record.values):
                raise InputError(f"{record.source}: selected column is outside the available fields.")
            selected.append(Record(record.source, tuple(record.values[c - 1] for c in options.columns)))
        dataset.records = selected
    return dataset
